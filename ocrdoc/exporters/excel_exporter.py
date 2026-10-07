from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ocrdoc.core.models import OCRResult, Table
from ocrdoc.exporters.base_exporter import BaseExporter


class ExcelExporter(BaseExporter):
    def export(self, result: OCRResult) -> bytes:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # type: ignore

        # Sheet "Full Text"
        ws_text = wb.create_sheet("Full Text")
        ws_text.column_dimensions["A"].width = 120
        for page in result.pages:
            ws_text.append([f"--- Trang {page.page_number} ---"])
            ws_text.cell(ws_text.max_row, 1).font = Font(bold=True)
            for block in page.text_blocks:
                if block.text.strip():
                    ws_text.append([block.text.strip()])
            ws_text.append([""])

        # Mỗi table → 1 sheet riêng
        table_num = 0
        for page in result.pages:
            for table in page.tables:
                table_num += 1
                sheet_name = f"Trang{page.page_number}_Bang{table_num}"[:31]
                ws = wb.create_sheet(sheet_name)
                self._write_table(ws, table)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _write_table(self, ws: Worksheet, table: Table) -> None:
        # Header row style
        header_fill = PatternFill("solid", fgColor="4472C4")
        header_font = Font(color="FFFFFF", bold=True)

        grid = table.to_2d()
        for r_idx, row in enumerate(grid):
            for c_idx, cell_text in enumerate(row):
                cell = ws.cell(row=r_idx + 1, column=c_idx + 1, value=cell_text)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if r_idx == 0:
                    cell.fill = header_fill
                    cell.font = header_font

        # Merged cells
        for tc in table.cells:
            if tc.row_span > 1 or tc.col_span > 1:
                ws.merge_cells(
                    start_row=tc.row + 1,
                    start_column=tc.col + 1,
                    end_row=tc.row + tc.row_span,
                    end_column=tc.col + tc.col_span,
                )

        # Auto-width (tối đa 50 chars)
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    @property
    def mime_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @property
    def file_extension(self) -> str:
        return "xlsx"
