from __future__ import annotations

import base64
import json

import openpyxl
import pytest

from ocrdoc.exporters import get_exporter
from ocrdoc.exporters.json_exporter import JSONExporter
from ocrdoc.exporters.markdown_exporter import MarkdownExporter
from ocrdoc.exporters.text_exporter import TextExporter
from ocrdoc.exporters.excel_exporter import ExcelExporter


class TestJSONExporter:
    def test_export_returns_valid_json(self, sample_ocr_result):
        exporter = JSONExporter()
        output = exporter.export(sample_ocr_result)
        data = json.loads(output)
        assert data["doc_type"] == "structured"
        assert len(data["pages"]) == 1

    def test_export_contains_text_blocks(self, sample_ocr_result):
        exporter = JSONExporter()
        output = exporter.export(sample_ocr_result)
        data = json.loads(output)
        blocks = data["pages"][0]["text_blocks"]
        assert len(blocks) == 3
        assert blocks[0]["text"] == "Công ty TNHH ABC"

    def test_export_contains_tables(self, sample_ocr_result):
        exporter = JSONExporter()
        output = exporter.export(sample_ocr_result)
        data = json.loads(output)
        tables = data["pages"][0]["tables"]
        assert len(tables) == 1
        assert tables[0]["rows"] == 3
        assert tables[0]["cols"] == 4

    def test_mime_type(self):
        assert JSONExporter().mime_type == "application/json"

    def test_file_extension(self):
        assert JSONExporter().file_extension == "json"


class TestMarkdownExporter:
    def test_export_returns_string(self, sample_ocr_result):
        exporter = MarkdownExporter()
        output = exporter.export(sample_ocr_result)
        assert isinstance(output, str)
        assert len(output) > 0

    def test_heading_rendered_correctly(self, sample_ocr_result):
        exporter = MarkdownExporter()
        output = exporter.export(sample_ocr_result)
        assert "## Công ty TNHH ABC" in output

    def test_table_rendered_as_pipe_table(self, sample_ocr_result):
        exporter = MarkdownExporter()
        output = exporter.export(sample_ocr_result)
        assert "| STT |" in output
        assert "| Tên hàng |" in output
        assert "|---|" in output or "|---" in output

    def test_mime_type(self):
        assert MarkdownExporter().mime_type == "text/markdown"


class TestTextExporter:
    def test_export_contains_all_text(self, sample_ocr_result):
        exporter = TextExporter()
        output = exporter.export(sample_ocr_result)
        assert "Công ty TNHH ABC" in output
        assert "Địa chỉ" in output
        assert "Sản phẩm A" in output

    def test_table_tab_separated(self, sample_ocr_result):
        exporter = TextExporter()
        output = exporter.export(sample_ocr_result)
        # Table cells nên được phân cách bằng tab
        assert "\t" in output

    def test_mime_type(self):
        assert TextExporter().mime_type == "text/plain"


class TestExcelExporter:
    def test_export_returns_bytes(self, sample_ocr_result):
        exporter = ExcelExporter()
        output = exporter.export(sample_ocr_result)
        assert isinstance(output, bytes)
        assert len(output) > 0

    def test_export_valid_xlsx(self, sample_ocr_result):
        exporter = ExcelExporter()
        output = exporter.export(sample_ocr_result)
        import io
        wb = openpyxl.load_workbook(io.BytesIO(output))
        assert "Full Text" in wb.sheetnames

    def test_table_sheet_created(self, sample_ocr_result):
        exporter = ExcelExporter()
        output = exporter.export(sample_ocr_result)
        import io
        wb = openpyxl.load_workbook(io.BytesIO(output))
        # Phải có sheet cho bảng
        assert len(wb.sheetnames) >= 2

    def test_table_data_correct(self, sample_ocr_result):
        exporter = ExcelExporter()
        output = exporter.export(sample_ocr_result)
        import io
        wb = openpyxl.load_workbook(io.BytesIO(output))
        # Tìm sheet có bảng
        table_sheets = [s for s in wb.sheetnames if s != "Full Text"]
        assert len(table_sheets) > 0
        ws = wb[table_sheets[0]]
        # Row đầu là header
        assert ws.cell(1, 1).value == "STT"
        assert ws.cell(1, 2).value == "Tên hàng"


class TestGetExporter:
    def test_get_json_exporter(self):
        exporter = get_exporter("json")
        assert isinstance(exporter, JSONExporter)

    def test_get_markdown_exporter(self):
        exporter = get_exporter("markdown")
        assert isinstance(exporter, MarkdownExporter)

    def test_get_md_alias(self):
        exporter = get_exporter("md")
        assert isinstance(exporter, MarkdownExporter)

    def test_get_txt_exporter(self):
        exporter = get_exporter("txt")
        assert isinstance(exporter, TextExporter)

    def test_get_xlsx_exporter(self):
        exporter = get_exporter("xlsx")
        assert isinstance(exporter, ExcelExporter)

    def test_invalid_format_raises(self):
        with pytest.raises(ValueError):
            get_exporter("docx")
